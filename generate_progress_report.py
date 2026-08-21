"""
generate_progress_report.py
===========================
Generates "PhishGuard AI — Project Progress Report I" in the exact format of
"Project Progress Report Format-1_26-27.docx" for college submission.

Formatting rules applied (from the format file):
    Page size      : A4
    Text font      : 12 point Times New Roman Regular
    Chapter heading: 14 point Times New Roman Bold
    Subtitle       : 12 point Times New Roman Bold
    Text Alignment : Justified
    Spacing        : 1.5 line spacing
    Margins        : Left 1.6 inch, Right/Top/Bottom 1.0 inch
    Figures/Tables : Centered; Table caption above, Figure caption below (size 12)
"""

import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------
FONT_NAME = "Times New Roman"
OUTPUT_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "PhishGuard_AI_Progress_Report_I.docx")

PROJECT_TITLE = "PhishGuard AI — An Intelligent Email Phishing Detection System Using Machine Learning"

DEPARTMENT_1 = "Department of Computer Science & Engineering"
COLLEGE_1 = "P. R. Pote Patil College of Engineering & Management, Amravati-444602 (M.S.)"

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def set_cell_text(cell, text, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, size=12, font_name=FONT_NAME):
    """Set a cell's text with formatting."""
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run(text)
    run.font.name = font_name
    run.font.size = Pt(size)
    run.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    # Vertical center
    tcPr = cell._tc.get_or_add_tcPr()
    vAlign = OxmlElement("w:vAlign")
    vAlign.set(qn("w:val"), "center")
    tcPr.append(vAlign)


def add_para(doc, text="", size=12, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
             line_spacing=1.5, space_after=6, font_name=FONT_NAME, underline=False):
    """Add a paragraph with the specified formatting."""
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run(text)
    run.font.name = font_name
    run.font.size = Pt(size)
    run.bold = bold
    run.underline = underline
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    return p


def add_heading(doc, text, size=14):
    """Add a chapter heading: 14pt Times New Roman Bold, centered."""
    return add_para(doc, text, size=size, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
                    line_spacing=1.5, space_after=10)


def add_subheading(doc, text, size=12):
    """Add a subtitle: 12pt Times New Roman Bold."""
    return add_para(doc, text, size=size, bold=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                    line_spacing=1.5, space_after=6)


def add_table_caption(doc, text):
    """Table caption ABOVE the table, centered, size 12."""
    return add_para(doc, text, size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
                    line_spacing=1.5, space_after=4)


def add_checkbox_line(doc, label):
    """Add a YES/NO checkbox line (using a bordered box) matching the format."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(4)

    run = p.add_run(label)
    run.font.name = FONT_NAME
    run.font.size = Pt(12)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)

    # YES checkbox
    run = p.add_run("    \u2610  YES")
    run.font.name = FONT_NAME
    run.font.size = Pt(12)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)

    # NO checkbox
    run = p.add_run("    \u2610  NO")
    run.font.name = FONT_NAME
    run.font.size = Pt(12)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    return p


def add_signature_line(doc, label, value=""):
    """Add a line for signature blocks."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(label)
    run.font.name = FONT_NAME
    run.font.size = Pt(12)
    run.bold = True
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    if value:
        run = p.add_run(value)
        run.font.name = FONT_NAME
        run.font.size = Pt(12)
        run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    return p


def set_table_borders(table):
    """Apply visible borders to the whole table."""
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "000000")
        borders.append(element)
    tblPr.append(borders)


# ---------------------------------------------------------------------------
# Document setup
# ---------------------------------------------------------------------------
doc = Document()

# Page size A4
section = doc.sections[0]
section.page_width = Inches(8.27)   # A4 width
section.page_height = Inches(11.69) # A4 height

# Margins: Left 1.6", Right/Top/Bottom 1.0"
section.left_margin = Inches(1.6)
section.right_margin = Inches(1.0)
section.top_margin = Inches(1.0)
section.bottom_margin = Inches(1.0)

# Set default style
style = doc.styles["Normal"]
style.font.name = FONT_NAME
style.font.size = Pt(12)
style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)

# ===========================================================================
# PAGE 1 — COVER PAGE (Report Header)
# ===========================================================================
for _ in range(2):
    add_para(doc, "", size=12, space_after=0)

add_para(doc, "Project Progress Report - I", size=14, bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.5, space_after=6)
add_para(doc, "on", size=12, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER,
         line_spacing=1.5, space_after=6)
add_para(doc, PROJECT_TITLE, size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
         line_spacing=1.5, space_after=6)

for _ in range(3):
    add_para(doc, "", size=12, space_after=0)

# Session / Group block
add_para(doc, "Session: 2026-27            Year/Semester: IVth / VIIth            Project Group No.: GA002",
         size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.5, space_after=6)
add_para(doc, "Project Guide Name: Prof. A. D. Chokhat            Project Leader Name: Mr. Abhay R. Kulkarni",
         size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.5, space_after=6)

for _ in range(2):
    add_para(doc, "", size=12, space_after=0)

add_para(doc, DEPARTMENT_1, size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
         line_spacing=1.5, space_after=4)
add_para(doc, COLLEGE_1, size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
         line_spacing=1.5, space_after=4)

doc.add_page_break()

# ===========================================================================
# PAGE 2 — TITLE PAGE (Submitted By)
# ===========================================================================
for _ in range(2):
    add_para(doc, "", size=12, space_after=0)

add_para(doc, "Project Progress Report - I", size=14, bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.5, space_after=6)
add_para(doc, "on", size=12, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER,
         line_spacing=1.5, space_after=6)
add_para(doc, PROJECT_TITLE, size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
         line_spacing=1.5, space_after=6)

for _ in range(2):
    add_para(doc, "", size=12, space_after=0)

add_para(doc, "Submitted By", size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
         line_spacing=1.5, space_after=10)

# Student names (two-column layout via a borderless table)
students_table = doc.add_table(rows=3, cols=2)
students_table.alignment = WD_TABLE_ALIGNMENT.CENTER
students_rows = [
    ("1. Name_of_Student-1", "3. Name_of_Student-3"),
    ("2. Name_of_Student-2", "4. Name_of_Student-4"),
    ("5. Name_of_Student-5", ""),
]
for r, (left, right) in enumerate(students_rows):
    set_cell_text(students_table.rows[r].cells[0], left, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(students_table.rows[r].cells[1], right, align=WD_ALIGN_PARAGRAPH.CENTER)

# Remove borders from the students table
borders_el = students_table._tbl.tblPr.find(qn("w:tblBorders"))
if borders_el is not None:
    students_table._tbl.tblPr.remove(borders_el)

for _ in range(3):
    add_para(doc, "", size=12, space_after=0)

add_para(doc, "<Guide Name> Guide", size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
         line_spacing=1.5, space_after=4)
add_para(doc, "Dr. V. B. Gadicha HOD, CSE", size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
         line_spacing=1.5, space_after=6)

for _ in range(3):
    add_para(doc, "", size=12, space_after=0)

add_para(doc, DEPARTMENT_1, size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
         line_spacing=1.5, space_after=4)
add_para(doc, COLLEGE_1, size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
         line_spacing=1.5, space_after=4)
add_para(doc, "(An Autonomous Institute)", size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
         line_spacing=1.5, space_after=4)
add_para(doc, "2026-2027", size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
         line_spacing=1.5, space_after=4)

doc.add_page_break()

# ===========================================================================
# PAGE 3 — TABLE OF CONTENTS
# ===========================================================================
for _ in range(2):
    add_para(doc, "", size=12, space_after=0)

add_heading(doc, "Table of Contents")

for _ in range(2):
    add_para(doc, "", size=12, space_after=0)

toc_data = [
    ("1", "Progress Summary", "1"),
    ("2", "Project Progress Details", "2"),
    ("3", "Work in Progress", "3"),
    ("4", "Project Planning", "4"),
    ("5", "Outcomes (Implications) Of Current Work", "4"),
    ("6", "Patent/Paper Publication/Copyright Certificate Status", "5"),
    ("7", "Conclusion", "5"),
]

toc_table = doc.add_table(rows=1, cols=3)
toc_table.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr = toc_table.rows[0].cells
for i, h in enumerate(["Sr. No.", "Contents", "Page No."]):
    set_cell_text(hdr[i], h, bold=True)
for sr, contents, page in toc_data:
    row = toc_table.add_row().cells
    set_cell_text(row[0], sr)
    set_cell_text(row[1], contents, align=WD_ALIGN_PARAGRAPH.LEFT)
    set_cell_text(row[2], page)
set_table_borders(toc_table)

doc.add_page_break()

# ===========================================================================
# PAGE 4 — OVERVIEW
# ===========================================================================
add_heading(doc, "Overview of Project Progress Report - I")

overview_data = [
    ("Number of Modules in your Projects:", "5"),
    ("Progress Report \u2013 I include how many modules:", "3"),
    ("Number of Objectives achieved in current progress report:", "4"),
]
overview_table = doc.add_table(rows=0, cols=2)
overview_table.alignment = WD_TABLE_ALIGNMENT.CENTER
for label, value in overview_data:
    row = overview_table.add_row().cells
    set_cell_text(row[0], label, align=WD_ALIGN_PARAGRAPH.LEFT)
    set_cell_text(row[1], value)
set_table_borders(overview_table)

add_para(doc, "", size=12, space_after=6)

add_subheading(doc, "Progress Summary")
add_para(doc,
         "[Give the following details in short paragraph or in bullet points]\n"
         "PhishGuard AI is a full-stack, machine-learning-powered email phishing detection platform. "
         "It combines a FastAPI backend, a React dashboard, and a scikit-learn (TF-IDF + Logistic Regression) "
         "inference engine to scan live emails in real time and assign a transparent 0\u2013100 risk score "
         "with explainable, human-readable reasons. The project is being developed as a final-year major "
         "project with the following modules and progress highlights:", size=12,
         align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_spacing=1.5, space_after=6)

add_para(doc,
         "\u2022  Module 1 (ML Model Training & Dataset Preparation) \u2013 Completed. "
         "A balanced dataset of 5,000 real-world emails (Enron Spam corpus from Hugging Face) was prepared and "
         "used to train a TF-IDF (1\u20132 n-grams) + handcrafted-feature Logistic Regression model tuned with "
         "GridSearchCV (5-fold stratified CV). Achieved Accuracy 98%, Precision 97.24%, Recall 98.8%, "
         "F1-Score 98.02%, ROC-AUC 99.58%.", size=12,
         align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_spacing=1.5, space_after=4)

add_para(doc,
         "\u2022  Module 2 (Backend API Development) \u2013 Completed. "
         "A production-grade FastAPI backend with JWT authentication (access + refresh tokens), "
         "rate limiting, structured logging, security headers, Alembic migrations, SQLite/PostgreSQL "
         "persistence, health/readiness endpoints, and audit logging.", size=12,
         align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_spacing=1.5, space_after=4)

add_para(doc,
         "\u2022  Module 3 (Frontend Dashboard Development) \u2013 Completed. "
         "A React 19 + Vite + Tailwind dashboard with a threat inbox, email detail view, explainable "
         "AI reasoning, System Analytics (Recharts), Raw Database view, Settings, and Login flow.", size=12,
         align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_spacing=1.5, space_after=4)

add_para(doc,
         "\u2022  Module 4 (Real-Time Threat Monitoring & WebSockets) \u2013 In Progress. "
         "Multi-provider IMAP daemon (Gmail, Outlook, Yahoo, Hotmail, Zoho, custom) with session "
         "persistence/auto-resume, manual refresh, and a real-time WebSocket feed with exponential backoff "
         "reconnection.", size=12,
         align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_spacing=1.5, space_after=4)

add_para(doc,
         "\u2022  Module 5 (Testing, Deployment & Documentation) \u2013 In Progress. "
         "44 backend pytest cases pass; frontend build succeeds; Docker Compose (PostgreSQL + backend + "
         "frontend/nginx) configuration is ready; deployment checklist and README are maintained.", size=12,
         align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_spacing=1.5, space_after=6)

add_subheading(doc, "List of Project Modules / Sections / Phases:")
add_para(doc,
         "1) ML Model Training & Dataset Preparation\n"
         "2) Backend API Development (FastAPI)\n"
         "3) Frontend Dashboard Development (React)\n"
         "4) Real-Time Threat Monitoring & WebSockets\n"
         "5) Testing, Deployment & Documentation", size=12,
         align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_spacing=1.5, space_after=6)

add_subheading(doc, "Status / Progress of Project Modules Completion")
add_para(doc,
         "Modules 1\u20133 are fully completed and verified. Module 4 is functionally complete "
         "(IMAP connect/refresh/disconnect and WebSocket alerts) and is under final testing. "
         "Module 5 is in progress \u2013 automated tests and Docker deployment are being finalized.", size=12,
         align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_spacing=1.5, space_after=6)

add_subheading(doc, "Observations / Inferences / Remarks (if any)")
add_para(doc,
         "The blended scoring (60% ML + 40% heuristics) with a trusted-domain allowlist effectively "
         "eliminates false positives on legitimate emails (Google, Amazon, Microsoft, etc.), while "
         "correctly flagging typosquatting domains, urgent-action requests, suspicious URLs, and "
         "PII requests. The model artifacts (model.pkl, vectorizer.pkl, metadata.json) load once at "
         "import time for fast inference.", size=12,
         align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_spacing=1.5, space_after=6)

add_subheading(doc, "Module-wise Project Progress Timeline")

timeline_data = [
    ("Module 1", "Completed", "01/08/2026"),
    ("Module 2", "Completed", "15/08/2026"),
    ("Module 3", "Completed", "25/08/2026"),
    ("Module 4", "In Progress", "30/09/2026"),
    ("Module 5", "In Progress", "31/10/2026"),
]

add_table_caption(doc, "Table 1.1: Module-wise Project Progress Timeline")
timeline_table = doc.add_table(rows=1, cols=3)
timeline_table.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr = timeline_table.rows[0].cells
for i, h in enumerate(["Module", "Status of Module", "Completion Date"]):
    set_cell_text(hdr[i], h, bold=True)
for mod, status, date in timeline_data:
    row = timeline_table.add_row().cells
    set_cell_text(row[0], mod)
    set_cell_text(row[1], status)
    set_cell_text(row[2], date)
set_table_borders(timeline_table)

doc.add_page_break()

# ===========================================================================
# PAGE 5 — DETAILS OF PROJECT PROGRESS REPORT-II
# ===========================================================================
add_heading(doc, "Details of Project Progress Report-II")

add_subheading(doc, "Modules planned for Progress Report-II")
rr2_data = [
    ("Real-Time Threat Monitoring & WebSockets", "30/09/2026"),
    ("Testing, Deployment & Documentation", "31/10/2026"),
]
add_table_caption(doc, "Table 2.1: Modules planned for Progress Report-II")
rr2_table = doc.add_table(rows=1, cols=2)
rr2_table.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr = rr2_table.rows[0].cells
for i, h in enumerate(["Name of Module", "Expected Completion Date"]):
    set_cell_text(hdr[i], h, bold=True)
for name, date in rr2_data:
    row = rr2_table.add_row().cells
    set_cell_text(row[0], name, align=WD_ALIGN_PARAGRAPH.LEFT)
    set_cell_text(row[1], date)
set_table_borders(rr2_table)

add_para(doc, "", size=12, space_after=6)

add_subheading(doc, "Required resources or support expected.")
add_para(doc,
         "\u2022  A live Gmail/Outlook/Yahoo test account (with app password) for validating the real-time "
         "IMAP monitoring module.\n"
         "\u2022  Google Cloud Console client (credentials.json) to demonstrate the Gmail OAuth integration.\n"
         "\u2022  A production host / container platform with sufficient memory for Docker-based deployment.\n"
         "\u2022  Guided review of the research paper draft and methodology.", size=12,
         align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_spacing=1.5, space_after=6)

add_subheading(doc, "Expected Outcomes (Implications) of Current Progress Work")
add_para(doc,
         "\u2022  A working full-stack phishing detection platform capable of scanning live inboxes in real time "
         "and classifying emails as legitimate, suspicious, or phishing with an explainable risk score.\n"
         "\u2022  High-accuracy ML model (98% accuracy, 98.8% recall) that minimizes false positives on "
         "trusted business emails.\n"
         "\u2022  A reusable, containerised deployment (Docker Compose) with secure JWT authentication, "
         "audit logging, and data export (CSV/JSON) for compliance.\n"
         "\u2022  Foundation for a research paper comparing phishing detection approaches and the blended "
         "ML + heuristic scoring methodology.", size=12,
         align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_spacing=1.5, space_after=6)

doc.add_page_break()

# ===========================================================================
# PAGE 6 — PATENT/PAPER PUBLICATION/COPYRIGHT CERTIFICATE STATUS
# ===========================================================================
add_heading(doc, "Patent/Paper Publication/Copyright Certificate status")

add_subheading(doc, "Have you started writing the research paper for Journal/Conference?")
add_para(doc, "\u2610  YES        \u2610  NO", size=12, align=WD_ALIGN_PARAGRAPH.LEFT, line_spacing=1.5)
add_para(doc, "", size=8, space_after=2)

add_subheading(doc, "Have you identified a suitable journal or conference for publication?")
add_para(doc, "\u2610  YES        \u2610  NO", size=12, align=WD_ALIGN_PARAGRAPH.LEFT, line_spacing=1.5)
add_para(doc, "", size=8, space_after=2)

add_subheading(doc, "Have you completed the literature survey?")
add_para(doc, "\u2610  YES        \u2610  NO", size=12, align=WD_ALIGN_PARAGRAPH.LEFT, line_spacing=1.5)
add_para(doc, "", size=8, space_after=2)

add_subheading(doc, "Have you prepared the problem statement based on the literature review?")
add_para(doc, "\u2610  YES        \u2610  NO", size=12, align=WD_ALIGN_PARAGRAPH.LEFT, line_spacing=1.5)
add_para(doc, "", size=8, space_after=2)

add_subheading(doc, "Has your guide reviewed the research paper draft?")
add_para(doc, "\u2610  YES        \u2610  NO", size=12, align=WD_ALIGN_PARAGRAPH.LEFT, line_spacing=1.5)
add_para(doc, "", size=8, space_after=2)

add_subheading(doc, "Have you finalized the research methodology?")
add_para(doc, "\u2610  YES        \u2610  NO", size=12, align=WD_ALIGN_PARAGRAPH.LEFT, line_spacing=1.5)
add_para(doc, "", size=8, space_after=2)

doc.add_page_break()

# ===========================================================================
# PAGE 7 — CONCLUSION
# ===========================================================================
add_heading(doc, "Conclusion")

add_para(doc,
         "In this first phase of the project, the core architecture of PhishGuard AI has been designed and "
         "implemented successfully. The machine-learning module was trained and tuned to deliver a 98% "
         "accuracy and an ROC-AUC of 99.58%, and the complete backend API along with the frontend dashboard "
         "are functional and integrated. The real-time IMAP monitoring and WebSocket alerting modules are "
         "under final testing, and deployment documentation has been prepared. The project demonstrates that "
         "a blend of ML scoring and hand-crafted heuristics with a trusted-domain allowlist provides a "
         "practical, low-false-positive phishing defence system. Future work in Progress Report-II will "
         "finalise live-monitoring validation, automated testing, containerised deployment, and the "
         "research-paper write-up.", size=12,
         align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_spacing=1.5, space_after=6)

doc.add_page_break()

# ===========================================================================
# PAGE 8 — GUIDE'S REMARKS
# ===========================================================================
add_heading(doc, "Guide's Remarks")

add_signature_line(doc, "Comments on progress work : ")
add_para(doc, "", size=12, space_after=4)
add_para(doc, "", size=12, space_after=4)
add_para(doc, "", size=12, space_after=4)

add_signature_line(doc, "Suggestions : ")
add_para(doc, "", size=12, space_after=4)
add_para(doc, "", size=12, space_after=4)
add_para(doc, "", size=12, space_after=4)

add_para(doc, "", size=12, space_after=10)
add_signature_line(doc, "Guide's signature : ")
add_para(doc, "", size=12, space_after=10)
add_signature_line(doc, "Date: ")

doc.add_page_break()

# ===========================================================================
# PAGE 9 — GROUP MEMBERS
# ===========================================================================
add_signature_line(doc, "Group Name : ")
add_para(doc, "", size=12, space_after=6)

members_data = [
    ("1", "Name_of_Student-1", "", "", ""),
    ("2", "Name_of_Student-2", "", "", ""),
    ("3", "Name_of_Student-3", "", "", ""),
    ("4", "Name_of_Student-4", "", "", ""),
    ("5", "Name_of_Student-5", "", "", ""),
]

add_table_caption(doc, "Table 3.1: Project Group Members")
members_table = doc.add_table(rows=1, cols=6)
members_table.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr = members_table.rows[0].cells
for i, h in enumerate(["Sr. No.", "Name of group member", "Role in Project",
                       "Email id", "Contact No.", "Sign"]):
    set_cell_text(hdr[i], h, bold=True)
for row_data in members_data:
    row = members_table.add_row().cells
    for i, val in enumerate(row_data):
        set_cell_text(row[i], val, align=WD_ALIGN_PARAGRAPH.LEFT if i == 1 else WD_ALIGN_PARAGRAPH.CENTER)
set_table_borders(members_table)

doc.add_page_break()

# ===========================================================================
# PAGE 10 — REPORT FORMATTING INSTRUCTIONS
# ===========================================================================
add_heading(doc, "Report Formatting Instructions")

formatting_data = [
    ("Page size", "A4"),
    ("Text font", "12 point Times New Roman Regular"),
    ("Chapter heading", "Title text 14 point Times New Roman Bold, Subtitle: 12 point Times New Roman Bold."),
    ("Text Alignment", "Justified"),
    ("Spacing", "1.5 (Line Spacing)"),
    ("Margins", "Left: - 1.6 inch. Right, Top, Bottom: - 1.0 inch."),
    ("Figures", "Figures should be centered placed numbered as Chapter no.(dot) Figure number (e.g. Figure. 3.1). Figure Caption must be below figure and centered with font size 12."),
    ("Tables", "Tables should be centered placed numbered as Chapter no.(dot) Table number (e.g. Table 3.1). Tables should be numbered separately. Table Caption must be above table and centered with font size 12."),
]
fmt_table = doc.add_table(rows=0, cols=2)
fmt_table.alignment = WD_TABLE_ALIGNMENT.CENTER
for label, value in formatting_data:
    row = fmt_table.add_row().cells
    set_cell_text(row[0], label, align=WD_ALIGN_PARAGRAPH.LEFT, bold=True)
    set_cell_text(row[1], value, align=WD_ALIGN_PARAGRAPH.LEFT)
set_table_borders(fmt_table)

add_para(doc, "", size=12, space_after=6)
add_para(doc, "Note:- The Project Progress Report should not exceed more than 05 pages.",
         size=12, bold=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_spacing=1.5)

# ---------------------------------------------------------------------------
# Save
# ---------------------------------------------------------------------------
doc.save(OUTPUT_PATH)
print(f"[OK] Report generated: {OUTPUT_PATH}")

