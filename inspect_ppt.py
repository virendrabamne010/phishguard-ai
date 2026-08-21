from pptx import Presentation
import sys

def inspect_ppt(file_path):
    try:
        prs = Presentation(file_path)
        for i, slide in enumerate(prs.slides):
            print(f"--- Slide {i+1} ---")
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    print(f"Text: {repr(shape.text)}")
    except Exception as e:
        print(f"Error: {e}")

if __name__ == "__main__":
    inspect_ppt(r"c:\Users\Virendra Bamne\OneDrive\Desktop\Project Progress Presentation.pptx")
